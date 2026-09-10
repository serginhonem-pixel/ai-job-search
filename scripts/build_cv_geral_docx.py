# -*- coding: utf-8 -*-
"""Gera cv/main_geral.docx replicando o estilo visual do bertolini_wide.docx
(Calibri, cabecalhos centralizados em negrito, corpo justificado), mas sem
faturamento/numero de funcionarios das empresas e sem o pacote de
remuneracao (que eram especificos daquela candidatura)."""
import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

d = docx.Document()

style = d.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(10)


def add(text, bold=False, size=10, align='justify'):
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if align == 'center' else WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    r.font.name = 'Calibri'
    return p


def blank():
    d.add_paragraph()


def role_block(company_line, cargo, resp_lines):
    add(company_line)
    add('Cargo: ' + cargo)
    add('Principais responsabilidades:')
    for line in resp_lines:
        add(line)
    blank()


# Header
add('Sergio Betini', bold=True, size=14, align='center')
add('Colatina - ES, Brasil')
add('+55 (27) 99783-6020')
add('sergiobetinim@gmail.com')
add('linkedin.com/in/sergio-betini')
blank()

# Formação Acadêmica
add('Formação Acadêmica:', bold=True, size=14, align='center')
add('MBA em Estratégias de Negócios, Competitividade e Produtividade – UNOPAR – 2023/2024')
add('MBA em Gestão Operacional e Logística Empresarial – UNOPAR – 2024')
add('Bacharelado em Administração – UNOPAR – 2020/2023')
blank()

# Resumo
add('Resumo:', bold=True, size=14, align='center')
add('Coordenador de PCP com histórico consistente em Planejamento e Controle de Produção, Supply Chain e Otimização de Processos Industriais. Combina conhecimento de chão de fábrica com ferramentas próprias de automação e domínio de ERPs industriais (TOTVS Protheus, Sankhya, SAP, Primus, Sagram) para reduzir custos e eliminar rupturas na cadeia de suprimentos.')
add('Habilidades-chave: PCP e Planejamento da Produção (MRP, S&OP), Otimização de Processos Industriais, Supply Chain Management, Business Intelligence (Power BI), gestão de ERPs industriais, liderança de equipes.')
add('Criou um otimizador de corte (Cutting Stock) que reduziu em 20% a geração de sucata na Metalosa.')
add('Desenvolveu simulador de estoque com integração EDI a fornecedores siderúrgicos, eliminando 100% da ruptura por falta de previsão.')
add('Liderou planejamento da transição de make-to-order para make-to-stock no Grupo Bertolini, reduzindo o lead time médio de 70 para 14 dias em todo o Brasil.')
add('Palavras-chave: PCP, MRP, S&OP, Cutting Stock, Supply Chain, Power BI, EDI, TOTVS Protheus, Sankhya, SAP.')
add('Projeto próprio: QtdApp (JavaScript), aplicativo de contagem/controle de inventário com integração a Excel e leitura por QR Code, já comercializado (qtdapp.vercel.app).')
blank()

# Idiomas
add('Idiomas:', bold=True, size=14, align='center')
add('Português – nativo')
add('Inglês – nível B1/B2 (conversacional)')
blank()

# Certificações
add('Certificações:', bold=True, size=14, align='center')
add('Power BI – Alura')
add('Excel VBA – Alura')
add('Excel – Alura')
blank()

# Experiência Profissional
add('Experiência Profissional:', bold=True, size=14, align='center')

role_block(
    'Metalosa – (Out/2025 – Atual) – Indústria metálica/metalurgia (Colatina, ES)',
    'Coordenador de PCP',
    [
        'Planejar e controlar a produção com análise de capacidade, integrando PCP, produção, logística e compras.',
        'Criou otimizador de corte (Cutting Stock) para o processo de Slitter, reduzindo em 20% a geração de sucata (~R$ 34.000,00 em matéria-prima economizada).',
        'Desenvolveu simulador de estoque com integração EDI a fornecedores siderúrgicos (Usiminas), eliminando 100% da ruptura por falta de previsão.',
        'Implantou painel Comercial-PCP, reduzindo o tempo de resposta entre comercial e produção de 2 dias para 2 horas.',
        'Criou sistema de gestão de manutenção com abertura de OS pela produção, reduzindo o tempo médio de atendimento em 70%.',
        'Automatizou a atualização de painéis via exportação agendada do ERP, eliminando trabalho manual e entregando dados quase em tempo real.',
    ]
)

role_block(
    'Unimarka – (Mai/2025 – Out/2025) – Distribuição/logística (Colatina, ES)',
    'Supervisor de Logística',
    [
        'Supervisionar operação de CD com equipe de 40 colaboradores, mantendo o atendimento no prazo.',
        'Reduziu retrabalho de separação em 33% via balanceamento de equipe e monitoramento diário de KPIs.',
        'Fez a interface com expedição, faturamento e transporte, padronizando as rotinas operacionais do CD.',
        'Coordenou o picking e o monitoramento diário de KPIs operacionais (volume separado, tempo de separação, retrabalho).',
    ]
)

role_block(
    'Distribuidora Caite – (Jul/2024 – Mai/2025) – Distribuição (Colatina, ES)',
    'Coordenador de Frota e Manutenção',
    [
        'Responsável pela gestão da frota, garantindo eficiência operacional e manutenção preventiva e corretiva dos veículos.',
        'Responsável pelo transporte de mercadorias entre fábricas e revenda, priorizando pontualidade, segurança e qualidade no serviço.',
        'Implantação de controles de calibragem de pneus, planos de manutenção e calendários de higienização para melhorar a conservação dos veículos.',
        'Criação e implementação de checklists diários de inspeção de veículos, assegurando segurança e organização na saída de rota.',
        'Monitorou o consumo de combustível, definindo metas de economia e usando telemetria para decisões estratégicas; renegociou contratos de diesel antes de reajustes de mercado.',
        'Renegociação de fornecedores, reduzindo o custo do GLP em 34% (R$ 8,75 -> R$ 5,75/kg) e do sabão industrial em 78% (R$ 3.200 -> R$ 700/mês).',
        'Planejamento de treinamentos para motoristas, alinhando a equipe aos padrões de segurança e eficiência exigidos.',
    ]
)

role_block(
    'Emerick Stones – (Fev/2024 – Jul/2024) – Rochas ornamentais/indústria de pedras (Colatina, ES)',
    'Analista de PCP Sênior',
    [
        'Análise de custos produtivos, identificando oportunidades de redução de despesas.',
        'Redesenho da lógica de programação da fábrica, aumentando eficiência e reduzindo desperdício.',
        'Conduziu a implantação do módulo de produção do ERP Sagram sem interrupção da produção.',
        'Criou e acompanhou KPIs para suporte à tomada de decisão estratégica.',
        'Criou procedimentos e protocolos de planejamento de produção, programação de ordens de trabalho e gestão de estoques, adaptados ao setor de beneficiamento de rochas.',
    ]
)

role_block(
    'Grupo Bertolini – (Set/2015 – Fev/2024) – Indústria de transformação (móveis de aço) (Colatina, ES)',
    'Analista de PPCP (Jun/2023 – Fev/2024)',
    [
        'Liderou a transição de make-to-order para make-to-stock, reduzindo o lead time de 70 para 14 dias em todo o Brasil.',
        'Manteve aderência ao plano de produção na gestão de demandas comerciais.',
        'Facilitou reuniões de S&OP entre vendas, produção, logística e finanças, com relato direto ao CEO.',
        'Acompanhou a implantação de um sistema MES (Manufatura) como Key User.',
    ]
)

add('Cargo: Assistente de PCP (Ago/2022 – Jun/2023)')
add('Principais responsabilidades:')
for line in [
    'Criou e implantou o planejamento de expedição (logística, itinerários, gestão de recursos).',
    'Geriu faturamento e carregamento por rota, otimizando custos de transporte.',
    'Desenvolveu estratégias para otimizar o carregamento de rotas, reduzindo custos e melhorando a eficiência do transporte.',
]:
    add(line)
blank()

add('Cargo: Líder de Equipe (Set/2015 – Set/2021)')
add('Principais responsabilidades:')
for line in [
    'Liderou equipe de 40 pessoas no setor de transformação por 6 anos, com foco em segurança, qualidade e produtividade.',
    'Coordenou treinamentos e desenvolvimento de equipe, garantindo atualização de habilidades e engajamento dos funcionários.',
    'Colaborou com compras, controle de qualidade e manutenção para garantir a integração eficaz das operações.',
]:
    add(line)
blank()

# Projeto próprio
add('Projeto Próprio:', bold=True, size=14, align='center')
add('QtdApp (qtdapp.vercel.app): aplicativo em JavaScript para contagem e controle de inventário, com integração a planilhas Excel, etiquetas personalizadas e leitura por QR Code. Já comercializado.')
blank()

# Referências
add('Referências:', bold=True, size=14, align='center')
add('Disponíveis mediante solicitação.')

d.save('cv/main_geral.docx')
print('saved cv/main_geral.docx')
